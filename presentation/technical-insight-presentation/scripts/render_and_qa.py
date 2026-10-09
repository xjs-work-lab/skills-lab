#!/usr/bin/env python3
"""Render an editable PPTX into per-page PNG, a PDF and contact sheet; perform basic geometry, notes and text checks.
Usage: python render_and_qa.py foo.pptx --outdir ./preview
Dependencies: libreoffice, pdftoppm, optional python-pptx, PIL/Pillow.
This tool does NOT assert aesthetic or scientific truth or native Microsoft PowerPoint fidelity.
"""
import argparse, json, shutil, subprocess, zipfile
from pathlib import Path

def run(cmd):
    v=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,text=True)
    if v.returncode: raise RuntimeError(f"command {cmd} failed: {v.stdout}\n{v.stderr}")
    return v.stdout

def main():
    ap=argparse.ArgumentParser();ap.add_argument('pptx');ap.add_argument('--outdir',default='preview');ap.add_argument('--dpi',type=int,default=144)
    args=ap.parse_args();src=Path(args.pptx).resolve();out=Path(args.outdir).resolve();out.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(src) as z:
        assert z.testzip() is None, 'Invalid PPTX ZIP'
        notes=[n for n in z.namelist() if n.startswith('ppt/notesSlides/notesSlide') and n.endswith('.xml')]
        slides=[n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]
    result=dict(pptx=str(src),notes_pages=len(notes),slide_parts=len(slides),overflow_shapes=[],pages=[],has_editable_text=False,has_native_shapes=False,manual_checks_required=['visual balance','text clipping in render','source/figure correctness','arrow meaning','PowerPoint font/link/room projection'])
    try:
        from pptx import Presentation
        pres=Presentation(src);result['page_count']=len(pres.slides)
        for i,s in enumerate(pres.slides,1):
            texts=0;shapes=0
            for sh in s.shapes:
                if sh.has_text_frame and sh.text.strip():texts+=1
                if sh.shape_type!=13:shapes+=1
                if sh.left<0 or sh.top<0 or sh.left+sh.width>pres.slide_width+3000 or sh.top+sh.height>pres.slide_height+3000:
                    result['overflow_shapes'].append(dict(page=i,name=sh.name,left=int(sh.left),top=int(sh.top)))
            if texts:result['has_editable_text']=True
            if shapes:result['has_native_shapes']=True
            result['pages'].append(dict(page=i,editable_text_shapes=texts,all_shapes=len(s.shapes),speaker_notes=bool(s.has_notes_slide)))
    except ImportError:
        result['python_pptx_available']=False
    soffice=shutil.which('soffice') or shutil.which('libreoffice');pdftoppm=shutil.which('pdftoppm')
    if soffice and pdftoppm:
        run([soffice,'-env:UserInstallation=file:///tmp/lo-insight-render-'+str(src.stat().st_ino), '--headless','--convert-to','pdf','--outdir',str(out),str(src)])
        pdf=out/(src.stem+'.pdf')
        if not pdf.exists():raise RuntimeError('PDF conversion returned success but file missing')
        run([pdftoppm,'-f','1','-png','-r',str(args.dpi),str(pdf),str(out/'slide')])
        result['rendered_pngs']=[str(x) for x in sorted(out.glob('slide-*.png'))]
        result['pdf']=str(pdf)
        try:
            from PIL import Image,ImageOps,ImageDraw
            imgs=[Image.open(p).convert('RGB') for p in result['rendered_pngs']]
            width=780;height=int(imgs[0].height*width/imgs[0].width)
            pad=16;cols=2;rows=(len(imgs)+cols-1)//cols
            board=Image.new('RGB',(cols*width+(cols+1)*pad,rows*height+(rows+1)*pad),'#EEF1F5')
            for i,im in enumerate(imgs):
                im.thumbnail((width,height))
                board.paste(im,(pad+(i%cols)*(width+pad),pad+(i//cols)*(height+pad)))
            board.save(out/'contact-sheet.png')
            result['contact_sheet']=str(out/'contact-sheet.png')
        except ImportError:pass
    else:result['renderer_missing']={'soffice':not bool(soffice),'pdftoppm':not bool(pdftoppm)}
    (out/'QA.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in result.items() if k not in ('pages','rendered_pngs')},ensure_ascii=False))
    if result['overflow_shapes']:raise SystemExit('Native shape overflow detected; see QA.json')

if __name__=='__main__':main()
